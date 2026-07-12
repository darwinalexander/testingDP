# -*- coding: utf-8 -*-
"""Módulo de datos compartido (40 árboles: 36 Parque Central + 4 Avenida El Oro).
IDs estables por código (Plus Code); veredictos y variables de salud/fitosanitarias."""
import os, glob, openpyxl
from openpyxl.utils import column_index_from_string as CIX

ASSETS = os.path.join(os.path.dirname(__file__), "..", "assets")
XLSX = os.path.join(ASSETS, "Reporte_Arboles_40.xlsx")
LOGO_UNL = os.path.join(ASSETS, "LOGO_UNL.png")
LOGO_ARBOLEC = os.path.join(ASSETS, "ArbolEC.jpeg")
MAPA = os.path.join(ASSETS, "mapa_arboles.png")

# Ortofoto opcional de fondo (assets/ortofoto.* + assets/ortofoto_bounds.txt)
_orto = [f for f in sorted(glob.glob(os.path.join(ASSETS, "ortofoto.*"))) if not f.endswith(".txt")]
ORTOFOTO = _orto[0] if _orto else None
ORTOFOTO_BOUNDS = None
_bf = os.path.join(ASSETS, "ortofoto_bounds.txt")
if os.path.exists(_bf):
    with open(_bf) as fh:
        v = [float(x) for x in fh.read().replace("\n", ",").split(",") if x.strip()]
        if len(v) == 4: ORTOFOTO_BOUNDS = v

BASE_URL = "https://arbolec.unl.edu.ec/ec/Saraguro"
def tree_url(codigo): return "%s/%s" % (BASE_URL, codigo)

# ---- ID estable por código (28 árboles ya analizados) ----
KNOWN_IDS = {
 "67829QG6+VHVQ":"A01","67829QG6+VHX3":"A02","67829QG6+VJ9V":"A03","67829QG6+WG7X":"A04",
 "67829QG6+WGX6":"A05","67829QG6+WJG5":"A06","67829QG6+XF9F":"A07","67829QG6+XFQG":"A08",
 "67829QG6+XG9W":"A09","67829QG6+XJQ9":"A10","67829QG6+XMQ4":"A11","67829QH6+2HV2":"A12",
 "67829QH6+2P3P":"A13","67829QH6+2QJF":"A14","67829QH6+3G9H":"A15","67829QH6+3HRM":"A16",
 "67829QH6+3JFH":"A17","67829QH6+3JR5":"A18","67829QH6+3M33":"A19","67829QH6+3MH3":"A20",
 "67829QH6+4H2W":"A21","67829QH6+4HHM":"A22","67829QH6+4HVV":"A23","67829QH6+4J2X":"A24",
 "67829QG6+XWHF":"AV01","67829QG6+XWMQ":"AV02","67829QG6+XX3P":"AV03","67829QH6+2V5H":"AV04",
}

# ---- Veredictos por código (28 previos + 12 nuevos) ----
VERDICTS = {
 "67829QG6+VHVQ": ("CONSERVAR-INT","Riesgo total alto asociado a fuste inclinado y daños mecánicos (heridas, ramas quebradas, poda inadecuada), no a inestabilidad de anclaje; caída Media. Ejemplar pequeño (D=6,3 cm; 3,2 m). Poda de reequilibrio, tratamiento de heridas y manejo de epífitas."),
 "67829QG6+VHX3": ("CONSERVAR","Buen estado estructural (fuste, corteza y hojas buenos); riesgo y caída bajos. Retirar cuerdas del fuste y poda ligera de formación."),
 "67829QG6+VJ9V": ("CONSERVAR-INT","Riesgo medio por fuste torcido, copa irregular y cima mala; caída Media. Especie nativa de valor ornamental. Poda de saneamiento y formación; tratar el descortezamiento."),
 "67829QG6+WG7X": ("CONSERVAR-INT","Condición fitosanitaria deficiente (hongos, exudados, pústulas, pudrición de ramas, ápice muerto), pero riesgo de CAÍDA BAJO por porte moderado (9,1 m) y raíz buena. Poda sanitaria, tratamiento fungicida, manejo de epífitas y monitoreo."),
 "67829QG6+WGX6": ("CONSERVAR-INT","Riesgo alto por afección foliar y de cima (hojas y cima malas, ápice muerto) y fuste inclinado; caída Media. Árbol joven-pequeño (D=19,7 cm; 7,4 m), recuperable. Poda sanitaria, fertilización y tratamiento; monitorear evolución."),
 "67829QG6+WJG5": ("DERRIBO","Único ejemplar con riesgo de CAÍDA ALTO y recomendación de derribo en campo. Es el árbol de mayor porte del parque (D=118,2 cm; 27,1 m), con copa amplia y poco simétrica —gran brazo de palanca frente a los fuertes vientos de la zona— exudados de resina (indicador de estrés/patología del fuste) y ubicación junto a los baños (máxima concurrencia de personas). La probabilidad de fallo combinada con la alta exposición vuelve inaceptable el riesgo. Se recomienda DERRIBO técnico controlado, previa verificación instrumental (resistógrafo/tomografía) si se dispone del equipo."),
 "67829QG6+XF9F": ("CONSERVAR-INT","Caída y riesgo Medios, pero RAÍZ en MAL estado en un árbol alto (17,8 m): punto crítico de anclaje. Poda de aligeramiento/reducción de la vela para disminuir la carga de viento e inspección radicular; si se confirma pérdida de anclaje, reevaluar derribo. Monitoreo PRIORITARIO."),
 "67829QG6+XFQG": ("CONSERVAR","Estructura y anclaje buenos; riesgo y caída bajos. Retiro de clavos/alambres y mantenimiento."),
 "67829QG6+XG9W": ("CONSERVAR","Ejemplar alto (20,7 m) en buen estado general, raíz buena y caída baja. Mantenimiento menor: retiro de clavos/alambres y manejo de epífitas."),
 "67829QG6+XJQ9": ("CONSERVAR-INT","Mayor VALOR de riesgo del parque (2,97), pero originado por daño a infraestructura por raíces (Alto) e inclinación en espacio estrecho, NO por deterioro estructural: fuste, ramas, hojas y cima BUENOS. Especie NATIVA (nogal andino). Poda de reequilibrio y manejo de raíz (barreras / reparación de adoquín); no procede derribo."),
 "67829QG6+XMQ4": ("CONSERVAR","Buen estado; riesgo y caída bajos. Poda menor y tratamiento de heridas."),
 "67829QH6+2HV2": ("CONSERVAR-INT","Buen anclaje (raíz buena) y caída baja. Poda ligera, manejo de epífitas y retiro de objetos extraños."),
 "67829QH6+2P3P": ("CONSERVAR","Palma en buen estado; riesgo mínimo. Mantenimiento (retiro de hojas secas)."),
 "67829QH6+2QJF": ("CONSERVAR-INT","Riesgo alto por fuste MUY inclinado en espacio estrecho y daño de raíz a infraestructura (Alto); caída Media. Poda de reequilibrio de copa para reducir carga en el lado inclinado y VIGILANCIA de la inclinación; si progresa, reevaluar."),
 "67829QH6+3G9H": ("CONSERVAR","Ejemplar joven en buen estado; riesgo mínimo. Mantenimiento."),
 "67829QH6+3HRM": ("CONSERVAR","Buen estado general; riesgo mínimo. Mantenimiento."),
 "67829QH6+3JFH": ("CONSERVAR","Riesgo bajo pese a fuste torcido; caída baja. Poda de formación y tratamiento de heridas."),
 "67829QH6+3JR5": ("CONSERVAR-INT","Riesgo bajo; cima mala y presencia de hongos. Poda sanitaria y manejo."),
 "67829QH6+3M33": ("CONSERVAR","Riesgo bajo, buen anclaje. Poda menor y manejo de epífitas."),
 "67829QH6+3MH3": ("CONSERVAR-INT","Especie NATIVA (molle). Riesgo alto por fuste MUY inclinado; caída Media, estructura buena (fuste y corteza buenos). Poda de reequilibrio y VIGILANCIA de la inclinación."),
 "67829QH6+4H2W": ("CONSERVAR-INT","Riesgo alto por inclinación y copa irregular; caída Media, fuste bueno. Poda de reequilibrio y de saneamiento."),
 "67829QH6+4HHM": ("CONSERVAR","Palma en buen estado; riesgo mínimo. Mantenimiento."),
 "67829QH6+4HVV": ("CONSERVAR-INT","Fuste bueno y raíz buena → riesgo de CAÍDA no crítico (Media); el problema es la afectación a construcciones, circulación y raíces (Alto) por su ubicación en esquina/espacio estrecho con raíces descubiertas. Corresponde manejo de raíz (barreras, reparación de adoquín) y poda, NO derribo."),
 "67829QH6+4J2X": ("CONSERVAR-INT","Riesgo alto por fuste muy inclinado; caída Media, ejemplar pequeño (D=7,6 cm). Poda de reequilibrio y tratamiento de heridas/descortezamiento."),
 "67829QG6+XWHF": ("CONSERVAR-INT","Riesgo alto por fuste inclinado en espacio estrecho, tumores en el fuste e interferencia con redes eléctricas; caída Media. Estructura general aceptable (cima y hojas buenas). Poda de reequilibrio y aclareo, tratamiento de daños mecánicos y manejo de epífitas; vigilancia de los tumores."),
 "67829QG6+XWMQ": ("CONSERVAR-INT","Especie NATIVA (molle). Presenta PUDRICIÓN DEL TRONCO, fuste y corteza en mal estado y tumores: defectos estructurales que, pese a una caída calificada Media, obligan a una EVALUACIÓN INSTRUMENTAL del fuste (resistógrafo/tomografía). Poda de descarga; si se confirma pérdida de resistencia del tronco, reevaluar derribo. Es el ejemplar de la avenida que requiere mayor atención."),
 "67829QG6+XX3P": ("CONSERVAR-INT","Riesgo medio por fuste torcido en espacio estrecho y daño de raíz a infraestructura (Alto); caída Media, fuste bueno. Manejo de raíz (barreras / reparación), poda de reequilibrio y tratamiento del descortezamiento."),
 "67829QH6+2V5H": ("CONSERVAR-INT","Árbol VIEJO y de gran copa (15,8 m) con raíz, fuste y corteza en mal estado, tumores, exudados y raíces descubiertas; caída Media pero con múltiples indicadores de senescencia. Poda de reducción para disminuir la carga de viento, inspección radicular y del fuste, y coordinación con la empresa eléctrica por la alta interferencia con redes. Reevaluar derribo si progresa el deterioro estructural. Monitoreo PRIORITARIO."),
 # --- 12 nuevos (Parque Central) ---
 "67829QG6+VH5Q": ("CONSERVAR-INT","Riesgo y caída bajos, pero se registra PUDRICIÓN DEL TRONCO (incipiente) junto con fuste inclinado e interferencia media con redes eléctricas. Poda de reequilibrio, coordinación con la empresa eléctrica y EVALUACIÓN/monitoreo del fuste por la pudrición; reevaluar si el defecto progresa."),
 "67829QG6+VJWX": ("CONSERVAR","Palma joven en buen estado; riesgo mínimo. Mantenimiento (retiro de hojas secas)."),
 "67829QG6+WGVV": ("CONSERVAR","Álamo joven en buen estado; riesgo y caída bajos. Mantenimiento y poda de formación."),
 "67829QG6+WHQJ": ("CONSERVAR","Palma en buen estado; leve afectación a construcciones por la base. Mantenimiento y vigilancia del entorno de la raíz."),
 "67829QG6+WJV4": ("CONSERVAR-INT","Riesgo alto por fuste inclinado; caída Media, pero estructura buena (fuste, corteza y hojas buenos). Poda de reequilibrio y vigilancia de la inclinación."),
 "67829QG6+XMX2": ("CONSERVAR","Riesgo bajo pese a fuste torcido; ejemplar de gran diámetro en buen estado. Poda de formación y de saneamiento."),
 "67829QH6+2GWF": ("CONSERVAR-INT","Riesgo medio por fuste torcido y afectación media a la circulación y a la infraestructura por raíces; caída Media. Poda de reequilibrio y manejo de raíz."),
 "67829QH6+2JHC": ("CONSERVAR","Palma adulta de gran porte, raíz y estructura buenas; follaje algo reducido. Mantenimiento (retiro de hojas secas) y riego."),
 "67829QH6+2JJ8": ("CONSERVAR-INT","Riesgo bajo; ejemplar joven ligeramente inclinado. Poda de formación para corregir el porte."),
 "67829QH6+2JRQ": ("CONSERVAR","Buen estado estructural; riesgo mínimo. Poda de formación."),
 "67829QH6+2P8P": ("CONSERVAR","Ejemplar joven en buen estado; riesgo mínimo. Mantenimiento."),
 "67829QH6+3P39": ("CONSERVAR","Riesgo mínimo; joven, fuste torcido. Poda de formación."),
}

# ---- Rangos de columnas de salud/fitosanitario (por letra) ----
FITO_COMPONENTS = ["Raíz","Fuste","Corteza","Ramas","Hojas","Cima","Copa"]   # AK-AQ
RNG_FITO = (CIX("AK"), CIX("AQ"))
RNG_ENF = (CIX("AR"), CIX("BB"))       # Presencia de enfermedades
RNG_UBIC_ENF = (CIX("BC"), CIX("BH"))  # Ubicación de las enfermedades
RNG_PLAGA = (CIX("BI"), CIX("BS"))     # Presencia de plagas
RNG_UBIC_PLAGA = (CIX("BT"), CIX("BX"))# Ubicación de las plagas

wb = openpyxl.load_workbook(XLSX, data_only=True)
ws = wb["Reporte"]
_h = {c: ws.cell(row=1, column=c).value for c in range(1, ws.max_column+1)}
def _col(name):
    for c, v in _h.items():
        if v == name: return c

def _vals_in(r, rng):
    out = []
    for c in range(rng[0], rng[1]+1):
        v = ws.cell(row=r, column=c).value
        if v is not None and str(v).strip():
            out.append(str(v).strip())
    return out

PARK_SITES = {"Parque Central", "Parque Central Saraguro"}
trees = []
_new_park = 24  # siguiente ID nuevo de parque será A25
for r in range(2, ws.max_row+1):
    cod = ws.cell(row=r, column=_col("Código")).value
    if not cod: continue
    def g(name):
        c = _col(name); return ws.cell(row=r, column=c).value if c else None
    sit_raw = g("Sitio")
    is_park = sit_raw in PARK_SITES
    if cod in KNOWN_IDS:
        tid = KNOWN_IDS[cod]
    elif is_park:
        _new_park += 1; tid = "A%02d" % _new_park
    else:
        tid = "AV?%d" % r
    fito = {}
    for i, comp in enumerate(FITO_COMPONENTS):
        v = ws.cell(row=r, column=RNG_FITO[0]+i).value
        fito[comp] = (str(v).strip() if v is not None and str(v).strip() else "-")
    cat, fund = VERDICTS.get(cod, ("CONSERVAR", "Sin observaciones relevantes; mantenimiento ordinario."))
    trees.append(dict(
        row=r, id=tid, codigo=cod,
        src=("A" if is_park else "AV"),
        sitio=("Parque Central" if is_park else "Avenida El Oro (isleta)"),
        especie=g("Especie"), familia=g("Familia"), comun=g("Nombre Común"),
        ref=(str(g("Comentario")) if g("Comentario") else ""),
        lon=g("Longitud"), lat=g("Latitud"), cf=g("Circunferencia de fuste"), d=g("Diámetro"),
        ht=g("Altura total"), hc=g("Altura comercial"),
        copaNS=g("Tam. copa (NS)"), copaEW=g("Tam. copa (EW)"), follaje=g("Follaje"),
        madurez=g("Estado de maduréz"), rectitud=g("Rectitud de fuste"), espacio=g("Espacio de crecimiento"),
        riesgo=g("Riesgo total"), valriesgo=g("Valor Riesgo total"),
        caida=g("Caida"), afec_constr=g("Afectación a construcciones"),
        afec_circ=g("Inteferencia con la circulación"), afec_raiz=g("Daño a infraestructura por raíces"),
        afec_redes=g("Interferencia con redes aéreas (eléctricas)"),
        fito=fito, raiz=fito["Raíz"], fuste_est=fito["Fuste"],
        enfermedades=_vals_in(r, RNG_ENF), ubic_enf=_vals_in(r, RNG_UBIC_ENF),
        plagas=_vals_in(r, RNG_PLAGA), ubic_plaga=_vals_in(r, RNG_UBIC_PLAGA),
        cat=cat, fund=fund, url=tree_url(cod)))

# ordenar: parque (A01..A36) y luego avenida (AV01..AV04)
def _idkey(t):
    p = 0 if t["src"] == "A" else 1
    num = int("".join(ch for ch in t["id"] if ch.isdigit()) or 0)
    return (p, num)
trees.sort(key=_idkey)
parque = [t for t in trees if t["src"] == "A"]
avenida = [t for t in trees if t["src"] == "AV"]

# ---- Estadísticos ----
CAT_LABEL = {"DERRIBO":"Derribo","CONSERVAR-INT":"Conservar con intervención","CONSERVAR":"Conservar"}
n = len(trees); n_parque = len(parque); n_av = len(avenida)
species = {}; families = {}
for t in trees:
    species[t["especie"]] = species.get(t["especie"], 0) + 1
    families[t["familia"]] = families.get(t["familia"], 0) + 1
risk_counts = {"Alto":0,"Medio":0,"Bajo":0}; caida_counts = {"Alto":0,"Medio":0,"Bajo":0}
for t in trees:
    risk_counts[t["riesgo"]] = risk_counts.get(t["riesgo"],0)+1
    caida_counts[t["caida"]] = caida_counts.get(t["caida"],0)+1
cupres = [t for t in trees if t["familia"] == "Cupressaceae"]
n_derribo = sum(1 for t in trees if t["cat"]=="DERRIBO")
n_int = sum(1 for t in trees if t["cat"]=="CONSERVAR-INT")
n_cons = sum(1 for t in trees if t["cat"]=="CONSERVAR")
comun_by = {}; fam_by = {}
for t in trees:
    comun_by.setdefault(t["especie"], t["comun"]); fam_by.setdefault(t["especie"], t["familia"])
DERRIBO_TREE = next(t for t in trees if t["cat"]=="DERRIBO")

# ---- Agregados de salud ----
def fito_color_key(v):
    v = (v or "").strip().lower()
    if v.startswith("buen"): return "buena"          # Buena/Bueno/Buenas/Buenos
    if v.startswith("regular"): return "regular"      # Regular/Regulares
    if v.startswith("mal"): return "mala"             # Mala/Malo/Malas/Malos
    if "poco" in v: return "regular"                  # Copa: Poco simétrica
    if "simétr" in v or "simetr" in v: return "buena" # Copa: Simétrica
    if "irregular" in v or "asimétr" in v: return "mala"  # Copa: Irregular
    return "na"                                       # No visible / No / -
# conteo por componente y estado
fito_summary = {comp: {"Buena/Bueno":0,"Regular":0,"Mala/Malo":0,"No visible/NA":0} for comp in FITO_COMPONENTS}
for t in trees:
    for comp in FITO_COMPONENTS:
        k = fito_color_key(t["fito"][comp])
        fito_summary[comp][{"buena":"Buena/Bueno","regular":"Regular","mala":"Mala/Malo","na":"No visible/NA"}[k]] += 1
def _freq(field):
    d = {}
    for t in trees:
        for x in t[field]:
            d[x] = d.get(x, 0) + 1
    return dict(sorted(d.items(), key=lambda kv: -kv[1]))
enf_freq = _freq("enfermedades")
plaga_freq = _freq("plagas")

def fnum(v, dec=1):
    if v is None: return "-"
    try: return ("%.*f" % (dec, float(v))).replace(".", ",")
    except: return str(v)

FOTOS_DIR = os.path.join(ASSETS, "fotos")
def tree_photos(tid):
    out = []
    for i in (1, 2, 3):
        hit = None
        if os.path.isdir(FOTOS_DIR):
            m = sorted(glob.glob(os.path.join(FOTOS_DIR, "%s_%d.*" % (tid, i))))
            if m: hit = m[0]
        out.append(hit)
    return out

if __name__ == "__main__":
    print("Total:", n, "| Parque:", n_parque, "| Avenida:", n_av)
    print("Especies:", len(species), "| Familias:", len(families), "| Cipreses:", len(cupres))
    print("Riesgo:", risk_counts, "| Caída:", caida_counts)
    print("Derribo:", n_derribo, "| Conservar-int:", n_int, "| Conservar:", n_cons)
    print("Derribo tree:", DERRIBO_TREE["id"], DERRIBO_TREE["codigo"])
    print("Nuevos IDs:", [t["id"] for t in parque if t["codigo"] not in KNOWN_IDS])
    print("Enfermedades:", enf_freq)
    print("Plagas:", plaga_freq)
