# -*- coding: utf-8 -*-
"""Genera versiones PDF del informe técnico y del oficio (ReportLab)."""
import openpyxl
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import cm, mm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.platypus import (BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer,
                               Table, TableStyle, PageBreak, NextPageTemplate, KeepTogether)

XLSX = "/root/.claude/uploads/4220178d-38bf-5cec-aaf2-15ce148fe2a2/6d02b559-Reporte__rboles_20260711__GAD_de_Saraguro.xlsx"

VERDE = colors.HexColor("#1B5E20"); GRIS = colors.HexColor("#424242")
ROJO = colors.HexColor("#B71C1C"); NARANJA = colors.HexColor("#E65A00"); AZUL = colors.HexColor("#0D47A1")
HDR = colors.HexColor("#1B5E20"); ROJO_F = colors.HexColor("#F4C7C3"); NAR_F = colors.HexColor("#FCE5CD"); VER_F = colors.HexColor("#D9EAD3")

# ---- datos (idéntico al docx) ----
wb = openpyxl.load_workbook(XLSX, data_only=True); ws = wb["Reporte"]
headers = {c: ws.cell(row=1, column=c).value for c in range(1, ws.max_column+1)}
def col(name):
    for c,h in headers.items():
        if h==name: return c
trees=[]
for r in range(2, ws.max_row+1):
    def g(name):
        c=col(name); return ws.cell(row=r,column=c).value if c else None
    trees.append(dict(row=r, id="A%02d"%(r-1), codigo=g("Código"), especie=g("Especie"),
        familia=g("Familia"), comun=g("Nombre Común"), ref=(str(g("Comentario")) if g("Comentario") else ""),
        lon=g("Longitud"), lat=g("Latitud"), d=g("Diámetro"), ht=g("Altura total"),
        copaNS=g("Tam. copa (NS)"), copaEW=g("Tam. copa (EW)"), follaje=g("Follaje"),
        rectitud=g("Rectitud de fuste"), espacio=g("Espacio de crecimiento"),
        riesgo=g("Riesgo total"), valriesgo=g("Valor Riesgo total"),
        raiz=g("Raíz"), fuste_est=g("Fuste"), caida=g("Caida"), afec_raiz=g("Daño a infraestructura por raíces")))
exec(open("/home/user/testingDP/_verdictos.py").read())  # define VER
for t in trees: t["cat"],t["fund"]=VER[t["row"]]
CAT={"DERRIBO":"Derribo","CONSERVAR-INT":"Conservar con intervención","CONSERVAR":"Conservar"}
CATF={"DERRIBO":ROJO_F,"CONSERVAR-INT":NAR_F,"CONSERVAR":VER_F}
RC={"Alto":ROJO,"Medio":NARANJA,"Bajo":VERDE}
def fnum(v,dec=1):
    if v is None: return "-"
    try: return ("%.*f"%(dec,float(v))).replace(".",",")
    except: return str(v)

n=len(trees)
species={}; families={}
for t in trees:
    species[t["especie"]]=species.get(t["especie"],0)+1
    families[t["familia"]]=families.get(t["familia"],0)+1
risk={"Alto":0,"Medio":0,"Bajo":0}; caida={"Alto":0,"Medio":0,"Bajo":0}
for t in trees:
    risk[t["riesgo"]]+=1; caida[t["caida"]]+=1
cupres=[t for t in trees if t["familia"]=="Cupressaceae"]
n_der=sum(1 for t in trees if t["cat"]=="DERRIBO")
n_int=sum(1 for t in trees if t["cat"]=="CONSERVAR-INT")
n_con=sum(1 for t in trees if t["cat"]=="CONSERVAR")
comun_by={}; fam_by={}
for t in trees:
    comun_by.setdefault(t["especie"],t["comun"]); fam_by.setdefault(t["especie"],t["familia"])

# ---- estilos ----
ss=getSampleStyleSheet()
def S(name,**kw):
    base=dict(fontName="Helvetica",fontSize=10.5,leading=14,alignment=TA_JUSTIFY)
    base.update(kw); return ParagraphStyle(name,**base)
body=S("body"); bodyc=S("bodyc",alignment=TA_CENTER)
h1=S("h1",fontName="Helvetica-Bold",fontSize=15,textColor=VERDE,spaceBefore=10,spaceAfter=6,alignment=TA_LEFT,leading=18)
h2=S("h2",fontName="Helvetica-Bold",fontSize=12,textColor=VERDE,spaceBefore=8,spaceAfter=4,alignment=TA_LEFT,leading=15)
small=S("small",fontSize=8.5,leading=11); smallc=S("smallc",fontSize=8.5,leading=11,alignment=TA_CENTER)
cellst=S("cell",fontSize=7.6,leading=9,alignment=TA_LEFT)
cellc=S("cellc",fontSize=7.6,leading=9,alignment=TA_CENTER)
def bullet_items(items,st=body):
    from reportlab.platypus import ListFlowable, ListItem
    return ListFlowable([ListItem(Paragraph(x,st),leftIndent=10) for x in items],bulletType="bullet",start="•",leftIndent=14)

def P(t,st=body): return Paragraph(t,st)

# ---- documento informe (portrait con secciones landscape) ----
def build_report(path):
    doc=BaseDocTemplate(path,pagesize=A4,leftMargin=2*cm,rightMargin=2*cm,topMargin=1.8*cm,bottomMargin=1.6*cm)
    fw,fh=A4
    portrait_frame=Frame(2*cm,1.6*cm,fw-4*cm,fh-3.4*cm,id="p")
    lw,lh=landscape(A4)
    land_frame=Frame(1.3*cm,1.3*cm,lw-2.6*cm,lh-2.6*cm,id="l")
    def foot(canvas,d):
        canvas.saveState(); canvas.setFont("Helvetica",7.5); canvas.setFillColor(GRIS)
        canvas.drawString(2*cm,1*cm,"UNL · Laboratorio de Dendrocronología — Informe Arbolado Parque Central de Saraguro")
        canvas.drawRightString(d.pagesize[0]-2*cm,1*cm,"Pág. %d"%canvas.getPageNumber()); canvas.restoreState()
    doc.addPageTemplates([
        PageTemplate(id="portrait",frames=[portrait_frame],pagesize=A4,onPage=foot),
        PageTemplate(id="land",frames=[land_frame],pagesize=landscape(A4),onPage=foot)])
    E=[]
    # PORTADA
    E+=[Spacer(1,1.4*cm),
        P("UNIVERSIDAD NACIONAL DE LOJA",S("x",fontName="Helvetica-Bold",fontSize=14,textColor=VERDE,alignment=TA_CENTER)),
        P("Facultad Agropecuaria y de Recursos Naturales Renovables",S("x",fontSize=11,textColor=GRIS,alignment=TA_CENTER)),
        P("Laboratorio de Dendrocronología y Anatomía de la Madera",S("x",fontSize=11,textColor=GRIS,alignment=TA_CENTER)),
        Spacer(1,1.3*cm),
        P("INFORME TÉCNICO DE EVALUACIÓN DEL ARBOLADO URBANO",S("x",fontName="Helvetica-Bold",fontSize=19,textColor=VERDE,alignment=TA_CENTER,leading=23)),
        P("Parque Central del cantón Saraguro",S("x",fontName="Helvetica-Bold",fontSize=15,textColor=GRIS,alignment=TA_CENTER)),
        Spacer(1,0.3*cm),
        P("Diagnóstico dendrométrico, estado fitosanitario y evaluación del riesgo de caída, con veredicto técnico de conservación o derribo",S("x",fontSize=11,textColor=GRIS,alignment=TA_CENTER,leading=15)),
        Spacer(1,1.1*cm),
        P("Solicitado por el Gobierno Autónomo Descentralizado Municipal Intercultural de Saraguro",smallc),
        P("Oficio Nro. 0286-A-GADMIS (19/05/2026) — Autorización Rectorado UNL-R-2026-2083-M (28/05/2026)",smallc),
        P("Referencia: UNL-SG-2026-0421-EX",smallc),
        Spacer(1,1.4*cm),
        P("Responsable técnico: Ph.D. Darwin Alexander Pucha Cofrep",S("x",fontSize=11,fontName="Helvetica-Bold",alignment=TA_CENTER)),
        P("Responsable del Laboratorio de Dendrocronología — Director de la Maestría en Biodiversidad y Cambio Climático",smallc),
        P("Equipo de campo: Darwin Pucha · Ariel Arévalo",smallc),
        Spacer(1,0.3*cm),
        P('Plataforma de consulta pública: <a href="https://arbolec.unl.edu.ec/ec/Saraguro" color="blue">https://arbolec.unl.edu.ec/ec/Saraguro</a>',smallc),
        P("Loja, 11 de julio de 2026",S("x",fontSize=10.5,fontName="Helvetica-Bold",textColor=GRIS,alignment=TA_CENTER)),
        PageBreak()]
    # 1 ANTECEDENTES
    E+=[P("1. Antecedentes y justificación",h1),
        P("El Gobierno Autónomo Descentralizado Municipal Intercultural de Saraguro, mediante Oficio Nro. 0286-A-GADMIS del 19 de mayo de 2026, suscrito por el Lic. Segundo Abel Sarango Quizhpe, Alcalde del cantón, solicitó a la Universidad Nacional de Loja el apoyo técnico para la inspección, evaluación y emisión de un criterio especializado sobre el estado actual de los árboles del Parque Central de Saraguro."),
        P("El Municipio expuso su preocupación por: (i) el crecimiento radicular de varios árboles, que ha ocasionado afectaciones en verjas, adoquinado y estructuras de concreto de las parcelas ornamentales; y (ii) el estado fitosanitario y la estabilidad estructural de dichos árboles frente a las fuertes corrientes de viento de la zona, situación que podría representar un riesgo para la seguridad ciudadana. El oficio destaca la presencia histórica de cipreses (registros fotográficos de la década de 1960), de reconocido valor simbólico y patrimonial para la comunidad saragurense."),
        P("Mediante Memorando Nro. UNL-R-2026-2083-M del 28 de mayo de 2026, el Rector, Dr. Nikolay Aguirre Mendoza, autorizó la participación del Ph.D. Darwin Alexander Pucha Cofrep, Responsable del Laboratorio de Dendrocronología, para el desarrollo de las actividades requeridas."),
        P("El presente informe responde a esa solicitud entregando un criterio técnico OBJETIVO y verificable que sustente la decisión de conservar o derribar cada árbol, con énfasis en el riesgo de caída y en los cipreses señalados. La recomendación de derribo se emite con carácter restrictivo: solo cuando la evidencia demuestra un riesgo inaceptable para las personas o los bienes."),
        # 2 OBJETIVOS
        P("2. Objetivos",h1),
        bullet_items([
            "Realizar el inventario y diagnóstico dendrométrico y fitosanitario del arbolado del Parque Central de Saraguro.",
            "Evaluar el nivel de riesgo de cada árbol, con énfasis en el riesgo de caída frente a las condiciones de viento y de sitio.",
            "Analizar de forma individual los cipreses del parque, por su relevancia patrimonial y por ser objeto de la preocupación municipal.",
            "Emitir un veredicto técnico por árbol —conservación o derribo— con su respectivo fundamento."]),
        # 3 METODOLOGIA
        P("3. Metodología",h1),
        P('El levantamiento se realizó mediante inspección visual en campo (metodología tipo VTA, <i>Visual Tree Assessment</i>) y registro georreferenciado con GPS, sistematizado en la plataforma institucional ArboLEC (<a href="https://arbolec.unl.edu.ec/ec/Saraguro" color="blue">arbolec.unl.edu.ec/ec/Saraguro</a>), donde el inventario queda publicado para consulta pública.'),
        P("Por cada individuo se registraron: identificación taxonómica, ubicación (coordenadas y código Plus Code), variables dendrométricas (circunferencia y diámetro, altura total y comercial, dimensiones de copa, follaje), estado de madurez, rectitud de fuste y espacio de crecimiento; la condición estructural y sanitaria de raíz, fuste, corteza, ramas, hojas, cima y copa; enfermedades, plagas y daños físicos; y los factores de riesgo (caída, afectación a construcciones, interferencia con la circulación, daño a infraestructura por raíces e interferencia con redes aéreas)."),
        P("<b>Escala de riesgo.</b> Cada factor se califica en Bajo, Medio y Alto; el sistema integra un Riesgo total (categórico) y un Valor de riesgo total (índice ~1,0–3,0). Se distinguen dos conceptos clave:"),
        bullet_items([
            "<b>Riesgo total / valor de riesgo:</b> índice global que combina TODOS los factores (incluido el daño de raíces a la infraestructura). Un valor alto no implica, por sí solo, peligro de caída.",
            "<b>Riesgo de caída:</b> probabilidad de fallo estructural (volcamiento o rotura). Es el factor determinante para la seguridad y el criterio rector para una eventual recomendación de derribo."]),
        P("<b>Criterio de veredicto.</b> Se recomienda DERRIBO únicamente cuando concurre un riesgo de caída Alto (o fallo estructural irreversible) con exposición de personas o bienes, y cuando el defecto no es corregible por poda o tratamiento. En los demás casos se recomienda CONSERVACIÓN, con o sin intervención. Este criterio, deliberadamente conservador, protege la seguridad ciudadana y el patrimonio arbóreo del parque."),
        PageBreak()]
    # 4 RESULTADOS
    E+=[P("4. Resultados generales del inventario",h1),
        P("Se inventariaron %d árboles (Sección S1), de %d especies y %d familias. La familia Cupressaceae (cipreses) es la más representada, con %d individuos (%.0f%%), lo que confirma su carácter dominante e identitario. El rango dendrométrico va de %s a %s m de altura y de %s a %s cm de diámetro; el mayor es el ciprés común A06 (código %s), con 27,1 m y 118,2 cm."%(
            n,len(species),len(families),len(cupres),100*len(cupres)/n,
            fnum(min(float(t['ht']) for t in trees)),fnum(max(float(t['ht']) for t in trees)),
            fnum(min(float(t['d']) for t in trees)),fnum(max(float(t['d']) for t in trees)),trees[5]['codigo']))]
    # tabla especies
    sp_sorted=sorted(species.items(),key=lambda x:(-x[1],x[0]))
    data=[["Especie","Nombre común","Familia","N°"]]
    for sp,cnt in sp_sorted: data.append([sp,comun_by.get(sp,""),fam_by.get(sp,""),str(cnt)])
    tsp=Table([[P(c,cellc if j==3 else cellst) for j,c in enumerate(row)] for row in data],
              colWidths=[5.5*cm,4.5*cm,4*cm,1.3*cm])
    tsp.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),HDR),("TEXTCOLOR",(0,0),(-1,0),colors.white),
        ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("FONTSIZE",(0,0),(-1,0),8.5),
        ("GRID",(0,0),(-1,-1),0.4,colors.grey),("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F3F7F3")])]))
    E+=[Spacer(1,4),P("Composición por especie:",S("b",fontName="Helvetica-Bold",fontSize=10)),tsp,Spacer(1,8)]
    # tabla riesgo total
    interp={"Alto":"Requieren intervención y/o seguimiento","Medio":"Intervención preventiva / monitoreo","Bajo":"Mantenimiento ordinario"}
    d2=[["Nivel de riesgo total","N°","%","Interpretación"]]
    for lvl in ["Alto","Medio","Bajo"]: d2.append([lvl,str(risk[lvl]),"%.0f%%"%(100*risk[lvl]/n),interp[lvl]])
    t2=Table([[P(c,cellc if j in(0,1,2) else cellst) for j,c in enumerate(row)] for row in d2],
             colWidths=[4.5*cm,1.5*cm,1.5*cm,7.8*cm])
    stl=[("BACKGROUND",(0,0),(-1,0),HDR),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),
        ("FONTSIZE",(0,0),(-1,0),8.5),("GRID",(0,0),(-1,-1),0.4,colors.grey),("VALIGN",(0,0),(-1,-1),"MIDDLE")]
    for i,lvl in enumerate(["Alto","Medio","Bajo"],start=1): stl.append(("TEXTCOLOR",(0,i),(0,i),RC[lvl]))
    t2.setStyle(TableStyle(stl))
    E+=[P("Distribución del riesgo total:",S("b",fontName="Helvetica-Bold",fontSize=10)),t2,
        P("Nota: el riesgo total integra todos los factores, incluido el daño de raíces a la infraestructura; por ello varios árboles estructuralmente sanos figuran en riesgo «Alto» sin ser candidatos a derribo.",S("i",fontSize=8.5,textColor=GRIS,alignment=TA_JUSTIFY,leading=11)),
        PageBreak()]
    # 5 CAIDA
    d3=[["Riesgo de caída","N°","%"]]
    for lvl in ["Alto","Medio","Bajo"]: d3.append([lvl,str(caida[lvl]),"%.0f%%"%(100*caida[lvl]/n)])
    t3=Table([[P(c,cellc) for c in row] for row in d3],colWidths=[4.5*cm,2*cm,2*cm])
    stl3=[("BACKGROUND",(0,0),(-1,0),HDR),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),
        ("FONTSIZE",(0,0),(-1,0),8.5),("GRID",(0,0),(-1,-1),0.4,colors.grey),("VALIGN",(0,0),(-1,-1),"MIDDLE")]
    for i,lvl in enumerate(["Alto","Medio","Bajo"],start=1): stl3.append(("TEXTCOLOR",(0,i),(0,i),RC[lvl]))
    t3.setStyle(TableStyle(stl3))
    E+=[P("5. Análisis del riesgo de caída (énfasis)",h1),
        P("El riesgo de caída es el criterio central por su relación directa con la seguridad de las personas. Distribución obtenida:"),
        t3,Spacer(1,6),
        P("<b>Hallazgo principal.</b> De los %d árboles, UN (1) solo individuo presenta riesgo de caída ALTO: el ciprés común A06 (código %s, ref. «%s»), que es además el único con recomendación de DERRIBO en campo. Diez (10) árboles presentan caída Media —manejables con poda de reequilibrio, tratamiento y vigilancia— y trece (13) caída Baja."%(n,trees[5]['codigo'],trees[5]['ref'])),
        P("Este resultado es determinante: aunque el 33%% del arbolado figura en riesgo total «Alto», ese nivel se asocia mayoritariamente a fustes inclinados y a daños de raíces sobre el adoquinado y el concreto —el problema que motivó la solicitud— y NO a una probabilidad real de volcamiento. En seguridad por caída, el parque tiene un único punto crítico. La inclinación y el daño radicular, sin implicar caída inminente, sí exigen seguimiento: se establecen como árboles de VIGILANCIA los A07 (raíz en mal estado, 17,8 m), A14 y A20 (fustes muy inclinados)."),
        P("<b>Recomendación de método:</b> para A06, y ante dudas futuras sobre ejemplares de gran porte, se aconseja confirmar el diagnóstico con evaluación instrumental (resistógrafo o tomografía sónica) antes de ejecutar el derribo."),
        PageBreak()]
    # 6 CIPRESES
    E+=[P("6. Análisis individual de los cipreses",h1),
        P("Atendiendo a la preocupación del Municipio y al valor patrimonial de estas especies, se analiza cada uno de los %d cipreses (4 <i>Cupressus sempervirens</i> — ciprés vela; 3 <i>Hesperocyparis macrocarpa</i> — ciprés común). Ambas especies son introducidas (Mediterráneo y California, respectivamente); su valor es cultural, histórico y paisajístico —reconocido en este informe— más que estrictamente ecológico."%len(cupres))]
    for t in cupres:
        rows=[
            ["Código / Ref. campo","%s / %s"%(t["codigo"],t["ref"] or "s/n"),"Coordenadas","%s, %s"%(fnum(t["lat"],6),fnum(t["lon"],6))],
            ["Diámetro / Altura","%s cm / %s m"%(fnum(t["d"]),fnum(t["ht"])),"Copa (NS×EW)","%s × %s m"%(fnum(t["copaNS"]),fnum(t["copaEW"]))],
            ["Rectitud / Espacio","%s / %s"%(t["rectitud"],t["espacio"]),"Follaje","%s%%"%t["follaje"]],
            ["Raíz / Fuste (estado)","%s / %s"%(t["raiz"],t["fuste_est"]),"Riesgo total (valor)","%s (%s)"%(t["riesgo"],fnum(t["valriesgo"],2))],
            ["RIESGO DE CAÍDA",t["caida"],"Afect. infraestructura (raíz)",t["afec_raiz"]]]
        tb=Table([[P(a,S("k",fontName="Helvetica-Bold",fontSize=8.3,leading=10)),P(b,small),
                   P(c,S("k",fontName="Helvetica-Bold",fontSize=8.3,leading=10)),P(d,small)] for a,b,c,d in rows],
                 colWidths=[4*cm,4.5*cm,4.3*cm,4*cm])
        tst=[("GRID",(0,0),(-1,-1),0.4,colors.grey),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("BACKGROUND",(0,0),(0,-1),colors.HexColor("#EEF4EE")),("BACKGROUND",(2,0),(2,-1),colors.HexColor("#EEF4EE"))]
        tst.append(("TEXTCOLOR",(1,4),(1,4),RC.get(t["caida"],GRIS)))
        tst.append(("FONTNAME",(1,4),(1,4),"Helvetica-Bold"))
        tb.setStyle(TableStyle(tst))
        catcol=ROJO if t["cat"]=="DERRIBO" else (NARANJA if t["cat"]=="CONSERVAR-INT" else VERDE)
        ver=P('<b>Veredicto: <font color="%s">%s.</font></b> %s'%(catcol.hexval(),CAT[t["cat"]],t["fund"]),body)
        block=[P("Ciprés %s — %s (%s)"%(t["id"],t["especie"],t["comun"]),
                 S("hc",fontName="Helvetica-Bold",fontSize=11.5,textColor=catcol,spaceBefore=6,spaceAfter=3)),
               tb,Spacer(1,3),ver,Spacer(1,6)]
        E.append(KeepTogether(block))
    E+=[P("<b>Síntesis de los cipreses:</b> de los %d cipreses se recomienda el derribo de UNO (A06) por riesgo de caída alto y alta exposición; los %d restantes se CONSERVAN con poda sanitaria, manejo de epífitas, tratamiento de daños mecánicos, retiro de clavos/alambres y, en A07, monitoreo prioritario del anclaje. Se conserva así el conjunto patrimonial de cipreses, retirando solo el ejemplar que representa un peligro real."%(len(cupres),len(cupres)-1)),
        NextPageTemplate("land"),PageBreak()]
    # 7 TABLA VEREDICTO (landscape)
    E+=[P("7. Tabla resumen: conservación vs. derribo (todo el inventario)",h1),
        P("Código de color: rojo = Derribo; naranja = Conservar con intervención; verde = Conservar. De los %d árboles: %d derribo, %d conservación con intervención y %d conservación con mantenimiento ordinario."%(n,n_der,n_int,n_con),small)]
    head=["ID","Especie (común)","Ref.","D (cm)","H (m)","Riesgo total","Caída","VEREDICTO","Fundamento"]
    dd=[[P(x,S("hh",fontName="Helvetica-Bold",fontSize=8,textColor=colors.white,alignment=TA_CENTER)) for x in head]]
    for t in trees:
        dd.append([P(t["id"],cellc),P("%s (%s)"%(t["especie"],t["comun"]),cellst),P(t["ref"] or "-",cellc),
                   P(fnum(t["d"]),cellc),P(fnum(t["ht"]),cellc),P("%s (%s)"%(t["riesgo"],fnum(t["valriesgo"],2)),cellc),
                   P('<font color="%s">%s</font>'%(RC.get(t["caida"],GRIS).hexval(),t["caida"]),cellc),
                   P(CAT[t["cat"]],cellc),P(t["fund"],cellst)])
    tv=Table(dd,colWidths=[1*cm,4*cm,1.5*cm,1.2*cm,1.2*cm,2*cm,1.3*cm,2.7*cm,11.5*cm],repeatRows=1)
    tvs=[("BACKGROUND",(0,0),(-1,0),HDR),("GRID",(0,0),(-1,-1),0.4,colors.grey),("VALIGN",(0,0),(-1,-1),"MIDDLE"),
         ("TOPPADDING",(0,0),(-1,-1),2),("BOTTOMPADDING",(0,0),(-1,-1),2)]
    for i,t in enumerate(trees,start=1): tvs.append(("BACKGROUND",(0,i),(-1,i),CATF[t["cat"]]))
    tv.setStyle(TableStyle(tvs))
    E+=[Spacer(1,4),tv,NextPageTemplate("portrait"),PageBreak()]
    # 8 FOTOS
    E+=[P("8. Registro fotográfico por árbol",h1),
        P('Las fotografías de cada árbol están publicadas en la plataforma ArboLEC junto con su ficha completa; cada individuo se localiza por su código y coordenadas en <a href="https://arbolec.unl.edu.ec/ec/Saraguro" color="blue">arbolec.unl.edu.ec/ec/Saraguro</a>. Se reserva el espacio para la fotografía de cada árbol; las imágenes se insertarán en la versión final del expediente municipal.'),Spacer(1,4)]
    for t in trees:
        catcol=ROJO if t["cat"]=="DERRIBO" else (NARANJA if t["cat"]=="CONSERVAR-INT" else VERDE)
        ph=P("<br/><br/>[ FOTOGRAFÍA ]<br/><b>%s</b><br/><i>(insertar imagen)</i><br/><br/>"%t["id"],S("ph",fontSize=9,alignment=TA_CENTER,textColor=GRIS,leading=12))
        info=P("<b><font color='%s'>%s — %s (%s)</font></b><br/>Código: %s | Ref.: %s<br/>Coordenadas: %s, %s<br/>D=%s cm · H=%s m · Follaje %s%%<br/>Riesgo total: %s (%s) · Caída: <font color='%s'>%s</font><br/><b><font color='%s'>Veredicto: %s</font></b><br/><font color='blue' size=8>Ficha: arbolec.unl.edu.ec/ec/Saraguro</font>"%(
            VERDE.hexval(),t["id"],t["especie"],t["comun"],t["codigo"],t["ref"] or "s/n",fnum(t["lat"],6),fnum(t["lon"],6),
            fnum(t["d"]),fnum(t["ht"]),t["follaje"],t["riesgo"],fnum(t["valriesgo"],2),RC.get(t["caida"],GRIS).hexval(),t["caida"],
            catcol.hexval(),CAT[t["cat"]]),S("info",fontSize=9,leading=12,alignment=TA_LEFT))
        ft=Table([[ph,info]],colWidths=[6*cm,11*cm])
        ft.setStyle(TableStyle([("BOX",(0,0),(-1,-1),0.5,colors.grey),("INNERGRID",(0,0),(-1,-1),0.4,colors.grey),
            ("BACKGROUND",(0,0),(0,0),colors.HexColor("#F2F2F2")),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("LEFTPADDING",(1,0),(1,0),8)]))
        E+=[KeepTogether([ft,Spacer(1,6)])]
    # 9 CONCLUSIONES
    E+=[PageBreak(),P("9. Conclusiones y recomendaciones",h1),
        P("Conclusiones:",S("b",fontName="Helvetica-Bold",fontSize=10.5)),
        bullet_items([
            "El Parque Central alberga %d árboles de %d especies; la familia de los cipreses es la dominante y de mayor valor patrimonial."%(n,len(species)),
            "En seguridad, existe UN único punto crítico de riesgo de caída Alto: el árbol A06 (ciprés común, código %s, junto a los baños), el más grande del parque y único con recomendación de derribo en campo."%trees[5]["codigo"],
            "El resto del arbolado (%d árboles) se CONSERVA; su riesgo total «Alto» se explica mayoritariamente por daño de raíces a la infraestructura y por inclinación del fuste, condiciones manejables sin derribo."%(n-1),
            "El daño al adoquinado y al concreto —motivo central de la solicitud— se atiende con manejo de raíces (barreras, reparación) y no justifica, por sí mismo, talar árboles sanos."]),
        P("Recomendaciones:",S("b",fontName="Helvetica-Bold",fontSize=10.5)),
        bullet_items([
            "Ejecutar el DERRIBO técnico controlado del árbol A06, previa verificación instrumental si se dispone del equipo, con personal especializado y medidas de seguridad por su cercanía a los baños.",
            "Programar la CONSERVACIÓN con intervención de los demás árboles: poda sanitaria y de reequilibrio, tratamiento fitosanitario, manejo de epífitas y retiro de clavos/alambres.",
            "Establecer MONITOREO prioritario de A07 (raíz en mal estado, 17,8 m) y de los ejemplares muy inclinados (A14, A20); reevaluar si progresa la inclinación o se confirma pérdida de anclaje.",
            "Implementar manejo de raíces (barreras anti-raíz, reparación del adoquinado) en los árboles con afectación Alta a infraestructura (A10, A23 y otros cipreses), preservando el árbol.",
            "Considerar un plan de reposición y sucesión a mediano plazo, priorizando especies nativas (nogal andino, molle, arupo).",
            "Mantener actualizado el inventario en ArboLEC como herramienta de gestión y seguimiento."]),
        Spacer(1,1.2*cm),
        P("____________________________________",bodyc),
        P("Ph.D. Darwin Alexander Pucha Cofrep",S("x",fontName="Helvetica-Bold",fontSize=11,alignment=TA_CENTER)),
        P("Responsable del Laboratorio de Dendrocronología — Universidad Nacional de Loja",smallc),
        P("Loja, 11 de julio de 2026",smallc)]
    doc.build(E)
    print("PDF informe:",path)

build_report("/home/user/testingDP/Informe_Tecnico_Arbolado_Parque_Central_Saraguro.pdf")

# ---- OFICIO PDF ----
def build_oficio(path):
    doc=BaseDocTemplate(path,pagesize=A4,leftMargin=2.5*cm,rightMargin=2.5*cm,topMargin=2*cm,bottomMargin=1.8*cm)
    fw,fh=A4
    fr=Frame(2.5*cm,1.8*cm,fw-5*cm,fh-3.8*cm,id="f")
    doc.addPageTemplates([PageTemplate(id="p",frames=[fr])])
    E=[P("UNIVERSIDAD NACIONAL DE LOJA",S("x",fontName="Helvetica-Bold",fontSize=13,textColor=VERDE,alignment=TA_CENTER)),
       P("Laboratorio de Dendrocronología y Anatomía de la Madera",smallc),
       P("Facultad Agropecuaria y de Recursos Naturales Renovables",smallc),Spacer(1,10),
       P("Oficio Nro. UNL-LDEND-2026-001",S("r",fontSize=10.5,alignment=TA_RIGHT)),
       P("Loja, 11 de julio de 2026",S("r",fontSize=10.5,alignment=TA_RIGHT)),Spacer(1,8),
       P("Licenciado",body),P("<b>Segundo Abel Sarango Quizhpe</b>",body),
       P("<b>ALCALDE DEL GOBIERNO AUTÓNOMO DESCENTRALIZADO MUNICIPAL INTERCULTURAL DE SARAGURO</b>",S("b",fontName="Helvetica-Bold",fontSize=10.5,leading=13)),
       P("Presente.-",body),Spacer(1,8),
       P("<b>ASUNTO:</b> Remisión del informe técnico de evaluación del arbolado del Parque Central de Saraguro y criterio de conservación/derribo.",body),
       P("<i>Referencias: Oficio Nro. 0286-A-GADMIS (19/05/2026); Memorando Nro. UNL-R-2026-2083-M (28/05/2026); Ref. UNL-SG-2026-0421-EX.</i>",S("i",fontSize=9.5,textColor=GRIS,alignment=TA_JUSTIFY,leading=12)),Spacer(1,6),
       P("De mi consideración:",body),
       P("Reciba un cordial saludo. En atención a su Oficio Nro. 0286-A-GADMIS y a la autorización conferida por el señor Rector mediante Memorando Nro. UNL-R-2026-2083-M, me permito remitir el Informe Técnico de Evaluación del Arbolado Urbano del Parque Central del cantón Saraguro, con el criterio especializado solicitado sobre el estado actual de los árboles y su eventual conservación o derribo.",body),
       P('El estudio comprendió el inventario y diagnóstico dendrométrico, fitosanitario y de riesgo de 24 árboles, sistematizado y publicado en la plataforma institucional ArboLEC (<a href="https://arbolec.unl.edu.ec/ec/Saraguro" color="blue">arbolec.unl.edu.ec/ec/Saraguro</a>), disponible para consulta pública. La evaluación atendió de manera particular a los cipreses del parque y al riesgo de caída frente a las condiciones de viento de la zona.',body),
       P("<b>Principales resultados:</b>",body),
       bullet_items([
         "Se recomienda el DERRIBO de UN (1) solo árbol: el ciprés común (<i>Hesperocyparis macrocarpa</i>) ubicado junto a los baños, identificado como A06 (código 67829QG6+WJG5). Es el de mayor porte del parque (27,1 m; 118,2 cm de diámetro), el único con riesgo de caída ALTO, con exudados de resina y copa amplia y poco simétrica que, ante los fuertes vientos y su cercanía a una zona de alta concurrencia, representa un riesgo inaceptable para las personas.",
         "Se recomienda CONSERVAR los 23 árboles restantes, incluidos los demás cipreses patrimoniales, mediante manejo: poda sanitaria y de reequilibrio, tratamiento fitosanitario, manejo de raíces y monitoreo de los ejemplares inclinados o con anclaje comprometido.",
         "El daño al adoquinado y al concreto por raíces —que motivó su solicitud— se atiende con manejo de raíces (barreras y reparación) sin necesidad de talar árboles estructuralmente sanos."]),
       P("Debo señalar, con la objetividad que exige este pronunciamiento, que la recomendación de derribo se ha emitido con criterio restrictivo: únicamente cuando la evidencia técnica demuestra un riesgo real e irreversible para la seguridad ciudadana, protegiendo así tanto a la población como el patrimonio arbóreo del Parque Central. Para el árbol A06 se sugiere, de ser posible, una verificación instrumental previa (resistógrafo o tomografía sónica) y su derribo por personal especializado con las debidas medidas de seguridad.",body),
       P("El detalle por árbol, el fundamento de cada veredicto, la tabla resumen de conservación y derribo y el registro fotográfico constan en el informe adjunto. Quedo a disposición del Municipio para socializar los resultados y acompañar técnicamente las acciones que se deriven.",body),
       P("Con sentimientos de distinguida consideración.",body),Spacer(1,10),
       P("Atentamente,",body),Spacer(1,18),
       P("____________________________________",body),
       P("<b>Ph.D. Darwin Alexander Pucha Cofrep</b>",body),
       P("Responsable del Laboratorio de Dendrocronología",S("g",fontSize=10,textColor=GRIS)),
       P("Universidad Nacional de Loja",S("g",fontSize=10,textColor=GRIS)),Spacer(1,10),
       P("<i>Adjunto: Informe Técnico de Evaluación del Arbolado del Parque Central de Saraguro.</i>",S("i",fontSize=9.5,textColor=GRIS)),
       P("<i>Copia: Dr. Nikolay Aguirre Mendoza, Rector de la Universidad Nacional de Loja.</i>",S("i",fontSize=9.5,textColor=GRIS))]
    doc.build(E)
    print("PDF oficio:",path)

build_oficio("/home/user/testingDP/Oficio_Respuesta_GAD_Saraguro.pdf")
