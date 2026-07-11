# -*- coding: utf-8 -*-
"""Oficio de respuesta al GAD Municipal de Saraguro (Word)."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import _data as D

OUT=os.path.join(os.path.dirname(__file__),"..","Oficio_Respuesta_GAD_Saraguro.docx")
VERDE=RGBColor(0x1B,0x5E,0x20); GRIS=RGBColor(0x42,0x42,0x42)

doc=Document()
st=doc.styles["Normal"]; st.font.name="Calibri"; st.font.size=Pt(11)
sec=doc.sections[0]; sec.top_margin=Cm(1.8); sec.bottom_margin=Cm(2); sec.left_margin=Cm(2.5); sec.right_margin=Cm(2.5)

def P(text,size=11,bold=False,italic=False,align="justify",color=None,after=8,before=0):
    p=doc.add_paragraph()
    p.alignment={"justify":WD_ALIGN_PARAGRAPH.JUSTIFY,"left":WD_ALIGN_PARAGRAPH.LEFT,"center":WD_ALIGN_PARAGRAPH.CENTER,"right":WD_ALIGN_PARAGRAPH.RIGHT}[align]
    r=p.add_run(text); r.bold=bold; r.italic=italic; r.font.size=Pt(size)
    if color: r.font.color.rgb=color
    p.paragraph_format.space_after=Pt(after); p.paragraph_format.space_before=Pt(before); return p
def bullet(text,size=11):
    p=doc.add_paragraph(style="List Bullet"); r=p.add_run(text); r.font.size=Pt(size); return p

# Logos
lt=doc.add_table(rows=1,cols=2); lt.alignment=WD_TABLE_ALIGNMENT.CENTER
lc=lt.rows[0].cells
lc[0].paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.LEFT
lc[0].paragraphs[0].add_run().add_picture(D.LOGO_UNL,width=Cm(6.5))
lc[1].paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.RIGHT
lc[1].paragraphs[0].add_run().add_picture(D.LOGO_ARBOLEC,width=Cm(2.4))
P("",after=2)
P("UNIVERSIDAD NACIONAL DE LOJA",size=12.5,bold=True,align="center",color=VERDE,after=0)
P("Laboratorio de Dendrocronología y Anatomía de la Madera · Facultad Agropecuaria y de Recursos Naturales Renovables",size=9.5,align="center",color=GRIS,after=12)

P("Oficio Nro. UNL-LDEND-2026-001",size=10.5,align="right",after=0)
P("Loja, 11 de julio de 2026",size=10.5,align="right",after=12)
P("Licenciado",after=0)
P("Segundo Abel Sarango Quizhpe",bold=True,after=0)
P("ALCALDE DEL GOBIERNO AUTÓNOMO DESCENTRALIZADO MUNICIPAL INTERCULTURAL DE SARAGURO",bold=True,size=10.5,after=0)
P("Presente.-",after=12)
P("ASUNTO: Remisión del informe técnico de evaluación del arbolado del Parque Central de Saraguro y criterio de conservación/derribo.",bold=True,after=6)
P("Referencias: Oficio Nro. 0286-A-GADMIS (19/05/2026); Memorando Nro. UNL-R-2026-2083-M (28/05/2026); Ref. UNL-SG-2026-0421-EX.",size=10,italic=True,color=GRIS,after=12)
P("De mi consideración:",after=8)
P("Reciba un cordial saludo. En atención a su Oficio Nro. 0286-A-GADMIS y a la autorización conferida por el señor Rector mediante Memorando Nro. UNL-R-2026-2083-M, me permito remitir el Informe Técnico de Evaluación del Arbolado Urbano del Parque Central del cantón Saraguro, con el criterio especializado solicitado sobre el estado actual de los árboles y su eventual conservación o derribo.")
P("El estudio comprendió el inventario y diagnóstico dendrométrico, fitosanitario y de riesgo de %d árboles (%d en el Parque Central y %d en la isleta de la Avenida El Oro), sistematizado y publicado en la plataforma institucional ArboLEC (%s), disponible para consulta pública. La evaluación atendió de manera particular a los cipreses del parque y al riesgo de caída frente a las condiciones de viento de la zona."%(D.n,D.n_parque,D.n_av,D.BASE_URL))
P("Principales resultados:",bold=True,after=4)
bullet("Se recomienda el DERRIBO de UN (1) solo árbol: el ciprés común (Hesperocyparis macrocarpa) ubicado junto a los baños, identificado como A06 (código %s). Es el de mayor porte del parque (27,1 m; 118,2 cm de diámetro), el único con riesgo de caída ALTO, con exudados de resina y copa amplia y poco simétrica que, ante los fuertes vientos y su cercanía a una zona de alta concurrencia, representa un riesgo inaceptable para las personas."%D.DERRIBO_TREE["codigo"])
bullet("Se recomienda CONSERVAR los %d árboles restantes, incluidos los demás cipreses patrimoniales, mediante manejo: poda sanitaria y de reequilibrio, tratamiento fitosanitario, manejo de raíces y monitoreo de los ejemplares inclinados o con anclaje comprometido."%(D.n-1))
bullet("En la isleta de la Avenida El Oro, dos árboles requieren evaluación estructural y monitoreo prioritario (sin ser derribo): un molle con pudrición del tronco (AV02) y una acacia vieja con raíz y fuste en mal estado (AV04).")
bullet("El daño al adoquinado y al concreto por raíces —que motivó su solicitud— se atiende con manejo de raíces (barreras y reparación) sin necesidad de talar árboles estructuralmente sanos.")
P("Debo señalar, con la objetividad que exige este pronunciamiento, que la recomendación de derribo se ha emitido con criterio restrictivo: únicamente cuando la evidencia técnica demuestra un riesgo real e irreversible para la seguridad ciudadana, protegiendo así tanto a la población como el patrimonio arbóreo del Parque Central. Para el árbol A06 se sugiere, de ser posible, una verificación instrumental previa (resistógrafo o tomografía sónica) y su derribo por personal especializado con las debidas medidas de seguridad.")
P("El detalle por árbol, el fundamento de cada veredicto, el mapa de ubicación, la tabla resumen de conservación y derribo y el registro fotográfico constan en el informe adjunto. Quedo a disposición del Municipio para socializar los resultados y acompañar técnicamente las acciones que se deriven.")
P("Con sentimientos de distinguida consideración.",after=18)
P("Atentamente,",after=24)
P("____________________________________",align="left",after=0)
P("Ph.D. Darwin Alexander Pucha Cofrep",bold=True,after=0)
P("Responsable del Laboratorio de Dendrocronología",size=10.5,color=GRIS,after=0)
P("Universidad Nacional de Loja",size=10.5,color=GRIS,after=14)
P("Adjunto: Informe Técnico de Evaluación del Arbolado del Parque Central y Avenida El Oro de Saraguro.",size=10,italic=True,color=GRIS,after=0)
P("Copia: Dr. Nikolay Aguirre Mendoza, Rector de la Universidad Nacional de Loja.",size=10,italic=True,color=GRIS,after=0)

doc.save(OUT)
print("Word oficio:",OUT)
